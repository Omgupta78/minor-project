"""
excel_report.py - the "full Excel document for the whole attendance".

Builds a single styled .xlsx workbook with four sheets:

  1. Summary          one row per student: sessions held, present, absent,
                      attendance %, and a DEFAULTER flag under 75%
  2. Attendance Grid  students down the side, every session across the top,
                      P / A / L in the cells - the register a teacher expects
  3. Detailed Records a flat, filterable log of every single mark, with the
                      recognition confidence and whether it was set by the
                      face engine or corrected by hand
  4. Session Log      one row per class photo taken: date, period, faces
                      detected, present/absent counts

Only openpyxl is required.
"""
from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import db

# ------------------------------------------------------------------- styles
INK = "2C2C2B"
MUTED = "7D7A75"
HEADER_BG = "1B4332"
BAND_BG = "F9F8F7"
BORDER_RGB = "E6E5E3"

GREEN_BG, GREEN_TXT = "E8F1EC", "46A171"
RED_BG, RED_TXT = "FCE9E7", "E56458"
ORANGE_BG, ORANGE_TXT = "FBEBDE", "D5803B"
BLUE_BG, BLUE_TXT = "E5F2FC", "2783DE"

TITLE_FONT = Font(name="Arial", size=14, bold=True, color=INK)
SUB_FONT = Font(name="Arial", size=9, color=MUTED)
HEAD_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color=INK)
MONO_FONT = Font(name="Consolas", size=10, color=INK)

HEAD_FILL = PatternFill("solid", fgColor=HEADER_BG)
BAND_FILL = PatternFill("solid", fgColor=BAND_BG)

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
HEAD_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

_thin = Side(style="thin", color=BORDER_RGB)
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

STATUS_STYLE = {
    "present": (GREEN_BG, GREEN_TXT, "P"),
    "late": (ORANGE_BG, ORANGE_TXT, "L"),
    "absent": (RED_BG, RED_TXT, "A"),
}


def _title_block(ws: Worksheet, title: str, subtitle: str, width: int) -> int:
    """Write a title + subtitle at the top. Returns the next free row."""
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.cell(row=2, column=1, value=subtitle).font = SUB_FONT
    if width > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 16
    return 4


def _header_row(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = HEAD_ALIGN
        cell.border = BOX
    ws.row_dimensions[row].height = 28


def _autosize(ws: Worksheet, min_width: int = 8, max_width: int = 42) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            col = cell.column
            if col is None:
                continue
            longest = max(len(line) for line in str(cell.value).split("\n"))
            widths[col] = max(widths.get(col, 0), longest)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(
            max(width + 3, min_width), max_width
        )


def _period_label(session) -> str:
    return f"P{session['period']}"


# ------------------------------------------------------------------ sheet 1
def _summary_sheet(wb: Workbook, conn, class_id, start, end, scope: str) -> None:
    ws = wb.create_sheet("Summary")
    rows = db.attendance_summary(conn, class_id, start, end)
    headers = [
        "Roll No",
        "Student Name",
        "Class",
        "Subject",
        "Sessions Held",
        "Present",
        "Absent",
        "Attendance %",
        "Status",
    ]
    row = _title_block(ws, "Attendance Summary", scope, len(headers))
    _header_row(ws, row, headers)
    first_data = row + 1

    for offset, item in enumerate(rows):
        r = first_data + offset
        values = [
            item["roll_no"],
            item["name"],
            item.get("class_name") or "-",
            item.get("subject") or "-",
            item["held"],
            item["present"],
            item["absent"],
            item["percent"] / 100.0,
            "DEFAULTER" if item["defaulter"] else "OK",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = MONO_FONT if col in (1, 5, 6, 7, 8) else BODY_FONT
            cell.alignment = LEFT if col in (2, 3, 4) else CENTER
            cell.border = BOX
            if offset % 2 == 1:
                cell.fill = BAND_FILL
        ws.cell(row=r, column=8).number_format = "0.0%"

        pct_cell = ws.cell(row=r, column=8)
        status_cell = ws.cell(row=r, column=9)
        if item["held"] == 0:
            tint, text = BLUE_BG, BLUE_TXT
        elif item["defaulter"]:
            tint, text = RED_BG, RED_TXT
        elif item["percent"] < 85:
            tint, text = ORANGE_BG, ORANGE_TXT
        else:
            tint, text = GREEN_BG, GREEN_TXT
        for cell in (pct_cell, status_cell):
            cell.fill = PatternFill("solid", fgColor=tint)
            cell.font = Font(name="Arial", size=10, bold=True, color=text)

    if not rows:
        ws.cell(row=first_data, column=1, value="No students enrolled yet.").font = BODY_FONT

    # Totals
    if rows:
        total_row = first_data + len(rows) + 1
        ws.cell(row=total_row, column=1, value="TOTAL").font = Font(
            name="Arial", size=10, bold=True, color=INK
        )
        ws.cell(row=total_row, column=2, value=f"{len(rows)} students").font = BODY_FONT
        held = rows[0]["held"] if rows else 0
        ws.cell(row=total_row, column=5, value=held).font = MONO_FONT
        ws.cell(row=total_row, column=6, value=sum(r["present"] for r in rows)).font = MONO_FONT
        ws.cell(row=total_row, column=7, value=sum(r["absent"] for r in rows)).font = MONO_FONT
        avg = sum(r["percent"] for r in rows) / len(rows) / 100.0
        avg_cell = ws.cell(row=total_row, column=8, value=avg)
        avg_cell.number_format = "0.0%"
        avg_cell.font = Font(name="Arial", size=10, bold=True, color=INK)
        for col in range(1, len(headers) + 1):
            ws.cell(row=total_row, column=col).border = BOX

        ws.auto_filter.ref = (
            f"A{first_data - 1}:{get_column_letter(len(headers))}{first_data + len(rows) - 1}"
        )

    ws.freeze_panes = ws.cell(row=first_data, column=3)
    _autosize(ws)


# ------------------------------------------------------------------ sheet 2
def _grid_sheet(wb: Workbook, conn, class_id, start, end, scope: str) -> None:
    ws = wb.create_sheet("Attendance Grid")
    students, sessions, marks = db.attendance_grid(conn, class_id, start, end)

    headers = ["Roll No", "Student Name"] + [
        f"{s['date'][5:]}\n{_period_label(s)}" for s in sessions
    ] + ["Present", "Held", "%"]
    row = _title_block(ws, "Attendance Register", scope, max(len(headers), 3))
    _header_row(ws, row, headers)
    first_data = row + 1

    for offset, student in enumerate(students):
        r = first_data + offset
        roll = ws.cell(row=r, column=1, value=student["roll_no"])
        roll.font, roll.alignment, roll.border = MONO_FONT, CENTER, BOX
        name = ws.cell(row=r, column=2, value=student["name"])
        name.font, name.alignment, name.border = BODY_FONT, LEFT, BOX
        if offset % 2 == 1:
            roll.fill = name.fill = BAND_FILL

        present = 0
        for idx, sess in enumerate(sessions):
            status = marks.get((student["id"], sess["id"]), "")
            bg, fg, letter = STATUS_STYLE.get(status, (None, MUTED, "-"))
            cell = ws.cell(row=r, column=3 + idx, value=letter)
            cell.alignment, cell.border = CENTER, BOX
            cell.font = Font(name="Arial", size=10, bold=True, color=fg)
            if bg:
                cell.fill = PatternFill("solid", fgColor=bg)
            if status in ("present", "late"):
                present += 1

        held = len(sessions)
        base = 3 + held
        for col, value in ((base, present), (base + 1, held)):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font, cell.alignment, cell.border = MONO_FONT, CENTER, BOX
        pct = ws.cell(row=r, column=base + 2, value=(present / held) if held else 0)
        pct.number_format = "0.0%"
        pct.alignment, pct.border = CENTER, BOX
        low = held and (present / held) < 0.75
        pct.font = Font(name="Arial", size=10, bold=True, color=RED_TXT if low else GREEN_TXT)
        pct.fill = PatternFill("solid", fgColor=RED_BG if low else GREEN_BG)

    if not students:
        ws.cell(row=first_data, column=1, value="No students enrolled yet.").font = BODY_FONT
    elif not sessions:
        ws.cell(
            row=first_data + len(students) + 1,
            column=1,
            value="No sessions recorded in this range - take attendance to fill the register.",
        ).font = SUB_FONT

    # Legend
    legend_row = first_data + max(len(students), 1) + 2
    ws.cell(row=legend_row, column=1, value="Legend").font = Font(
        name="Arial", size=9, bold=True, color=MUTED
    )
    for idx, (label, (bg, fg, letter)) in enumerate(STATUS_STYLE.items()):
        cell = ws.cell(row=legend_row, column=2 + idx * 2, value=letter)
        cell.font = Font(name="Arial", size=9, bold=True, color=fg)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.alignment, cell.border = CENTER, BOX
        ws.cell(row=legend_row, column=3 + idx * 2, value=label.title()).font = SUB_FONT

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 26
    for idx in range(len(sessions)):
        ws.column_dimensions[get_column_letter(3 + idx)].width = 7
    for idx in range(3):
        ws.column_dimensions[get_column_letter(3 + len(sessions) + idx)].width = 9
    ws.freeze_panes = ws.cell(row=first_data, column=3)


# ------------------------------------------------------------------ sheet 3
def _detail_sheet(wb: Workbook, conn, class_id, start, end, scope: str) -> None:
    ws = wb.create_sheet("Detailed Records")
    rows = db.detailed_records(conn, class_id, start, end)
    headers = [
        "Date",
        "Period",
        "Class",
        "Subject",
        "Roll No",
        "Student Name",
        "Status",
        "Confidence",
        "Marked By",
        "Marked At",
    ]
    row = _title_block(ws, "Detailed Attendance Records", scope, len(headers))
    _header_row(ws, row, headers)
    first_data = row + 1

    for offset, rec in enumerate(rows):
        r = first_data + offset
        values = [
            rec["date"],
            f"P{rec['period']}",
            rec["class_name"],
            rec["subject"] or "-",
            rec["roll_no"],
            rec["name"],
            rec["status"].title(),
            rec["confidence"] if rec["confidence"] is not None else "",
            "Face engine" if rec["method"] == "face" else "Teacher",
            rec["marked_at"],
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = MONO_FONT if col in (1, 2, 5, 8, 10) else BODY_FONT
            cell.alignment = LEFT if col in (3, 4, 6, 9) else CENTER
            cell.border = BOX
            if offset % 2 == 1:
                cell.fill = BAND_FILL
        ws.cell(row=r, column=8).number_format = "0.0%"

        bg, fg, _ = STATUS_STYLE.get(rec["status"], (BLUE_BG, BLUE_TXT, ""))
        status_cell = ws.cell(row=r, column=7)
        status_cell.fill = PatternFill("solid", fgColor=bg)
        status_cell.font = Font(name="Arial", size=10, bold=True, color=fg)

    if not rows:
        ws.cell(row=first_data, column=1, value="No attendance recorded yet.").font = BODY_FONT
    else:
        ws.auto_filter.ref = (
            f"A{first_data - 1}:{get_column_letter(len(headers))}{first_data + len(rows) - 1}"
        )
    ws.freeze_panes = ws.cell(row=first_data, column=1)
    _autosize(ws)


# ------------------------------------------------------------------ sheet 4
def _session_sheet(wb: Workbook, conn, class_id, start, end, scope: str) -> None:
    ws = wb.create_sheet("Session Log")
    sessions = db.list_sessions(conn, class_id, start, end)
    headers = [
        "Date",
        "Period",
        "Class",
        "Subject",
        "Faces Detected",
        "Present",
        "Absent",
        "Taken By",
        "Recorded At",
    ]
    row = _title_block(ws, "Session Log", scope, len(headers))
    _header_row(ws, row, headers)
    first_data = row + 1

    for offset, sess in enumerate(sessions):
        r = first_data + offset
        values = [
            sess["date"],
            f"P{sess['period']}",
            sess["class_name"],
            sess["class_subject"] or "-",
            sess["total_faces"],
            sess["present_count"],
            sess["absent_count"],
            sess["taken_by"] or "-",
            sess["created_at"],
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = MONO_FONT if col in (1, 2, 5, 6, 7, 9) else BODY_FONT
            cell.alignment = LEFT if col in (3, 4, 8) else CENTER
            cell.border = BOX
            if offset % 2 == 1:
                cell.fill = BAND_FILL

    if not sessions:
        ws.cell(row=first_data, column=1, value="No sessions recorded yet.").font = BODY_FONT
    ws.freeze_panes = ws.cell(row=first_data, column=1)
    _autosize(ws)


# --------------------------------------------------------------------- API
def build_workbook(
    conn,
    class_id: Optional[int] = None,
    start: Optional[str] = None,
    end: Optional[str] = None,
) -> Workbook:
    """Build the complete attendance workbook."""
    klass = db.get_class(conn, class_id) if class_id else None
    parts = []
    if klass:
        label = klass["name"]
        if klass["subject"]:
            label += f" - {klass['subject']}"
        parts.append(label)
    else:
        parts.append("All classes")
    if start or end:
        parts.append(f"{start or 'start'} to {end or 'today'}")
    parts.append(f"generated {datetime.now().strftime('%d %b %Y, %H:%M')}")
    scope = "  |  ".join(parts)

    wb = Workbook()
    wb.remove(wb.active)
    _summary_sheet(wb, conn, class_id, start, end, scope)
    _grid_sheet(wb, conn, class_id, start, end, scope)
    _detail_sheet(wb, conn, class_id, start, end, scope)
    _session_sheet(wb, conn, class_id, start, end, scope)

    props = wb.properties
    props.title = "Attendance Report"
    props.creator = "Face Recognition Attendance System"
    props.created = datetime.now()
    return wb


def workbook_bytes(conn, class_id=None, start=None, end=None) -> io.BytesIO:
    buffer = io.BytesIO()
    build_workbook(conn, class_id, start, end).save(buffer)
    buffer.seek(0)
    return buffer


def suggested_filename(conn, class_id: Optional[int] = None) -> str:
    klass = db.get_class(conn, class_id) if class_id else None
    stem = "attendance"
    if klass:
        safe = "".join(ch if ch.isalnum() else "_" for ch in klass["name"]).strip("_")
        stem = f"attendance_{safe}" if safe else stem
    return f"{stem}_{date.today().isoformat()}.xlsx"
